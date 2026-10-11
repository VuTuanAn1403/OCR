"""Export the approved Markdown specification using bundled python-docx.

The Markdown remains the content authority. No empirical results are added.
"""
from pathlib import Path
import hashlib
import json
import re
import textwrap

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'PROJECT_SPECIFICATION_RESEARCH_V5.md'
OUTPUT = ROOT / 'PROJECT_SPECIFICATION_RESEARCH_V5.docx'
QA = ROOT / 'specifications' / 'word_qa'
QA.mkdir(parents=True, exist_ok=True)
source_text = SOURCE.read_text(encoding='utf-8')


def element(tag, **attributes):
    e = OxmlElement(tag)
    for key, value in attributes.items():
        e.set(qn(key), str(value))
    return e


def set_font(run, name='Calibri', size=None, bold=None, italic=None, color='000000'):
    run.font.name = name
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    props = run._r.get_or_add_rPr()
    props.append(element('w:rFonts', **{'w:ascii': name, 'w:hAnsi': name, 'w:eastAsia': name, 'w:cs': name}))


def hyperlink(paragraph, label, href, color='175C83', size=None):
    if not re.match(r'^[a-z]+://', href, re.I):
        href = (SOURCE.parent / href).resolve().as_uri()
    link = element('w:hyperlink', **{'r:id': paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)})
    run = element('w:r')
    props = element('w:rPr')
    props.append(element('w:color', **{'w:val': color}))
    props.append(element('w:u', **{'w:val': 'single'}))
    if size:
        props.append(element('w:sz', **{'w:val': str(int(size*2))}))
    run.append(props)
    t = element('w:t')
    t.text = label
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


INLINE = re.compile(r'(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)')


def inline(p, text, table=False, header=False):
    base_color = 'FFFFFF' if header else '000000'
    base_size = 10.5 if table else 11
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', chunk)
        if m:
            hyperlink(p, m[1], m[2], base_color if header else '175C83', base_size)
            continue
        if chunk.startswith('`') and chunk.endswith('`'):
            r = p.add_run(chunk[1:-1])
            set_font(r, 'Consolas', 9 if table else 9.5, bold=header, color=base_color)
        elif chunk.startswith('**') and chunk.endswith('**'):
            r = p.add_run(chunk[2:-2])
            set_font(r, size=base_size, bold=True, color=base_color)
        elif chunk.startswith('*') and chunk.endswith('*'):
            r = p.add_run(chunk[1:-1])
            set_font(r, size=base_size, italic=True, bold=header, color=base_color)
        else:
            r = p.add_run(chunk)
            set_font(r, size=base_size, bold=header, color=base_color)


def draw_architecture():
    im = Image.new('RGB', (1500, 1230), 'white')
    d = ImageDraw.Draw(im)
    font = ImageFont.truetype(r'C:\Windows\Fonts\calibri.ttf', 36)
    small = ImageFont.truetype(r'C:\Windows\Fonts\calibri.ttf', 32)
    boxes = {}

    def box(key, label, xy, width=630, height=120):
        x, y = xy
        boxes[key] = (x, y, x+width, y+height)
        d.rounded_rectangle(boxes[key], radius=12, fill='#F5F7FA', outline='#506172', width=2)
        lines = textwrap.wrap(label, width=31 if width>400 else 19)
        spacing = 39 if width>400 else 36
        yy = y+(height-len(lines)*spacing)/2
        for line in lines:
            f = font if width>400 else small
            bb = d.textbbox((0,0), line, font=f)
            d.text((x+(width-(bb[2]-bb[0]))/2, yy), line, font=f, fill='#000000')
            yy += spacing

    def arrow(points):
        d.line(points, fill='#506172', width=3)
        x0,y0 = points[-2]
        x,y = points[-1]
        if abs(y-y0)>=abs(x-x0):
            s = 1 if y>y0 else -1
            d.polygon([(x,y),(x-10,y-15*s),(x+10,y-15*s)], fill='#506172')
        else:
            s = 1 if x>x0 else -1
            d.polygon([(x,y),(x-15*s,y-10),(x-15*s,y+10)], fill='#506172')

    labels = [
        ('A','PDF và cấu hình đã khóa',60,35),
        ('B','Kiểm tra đầu vào và định danh nguồn',60,195),
        ('C','Inspect trang và định tuyến theo vùng',60,355),
        ('F','IR có provenance và tọa độ',60,710),
        ('G','Layout bảng thứ tự đọc và kiểm tra lỗi',60,870),
        ('M','GT độc lập đã khóa',60,1070),
        ('H','Retry vùng trong ngân sách',820,35),
        ('I','Export Markdown JSON và ảnh',820,195),
        ('J','Kiểm tra artifact và commit đầu ra',820,355),
        ('K','Đóng thời gian end to end và manifest',820,515),
        ('L','Evaluator đọc Markdown cuối',820,750),
        ('N','Counts CI báo cáo và claim evidence',820,990),
    ]
    for key,label,x,y in labels:
        box(key,label,(x,y))
    box('D','Trích xuất native',(60,530),290,130)
    box('E','Render và OCR CPU',(400,530),290,130)
    for a,b in [('A','B'),('B','C'),('F','G'),('H','I'),('I','J'),('J','K'),('K','L'),('L','N')]:
        aa,bb=boxes[a],boxes[b]
        arrow([((aa[0]+aa[2])/2,aa[3]),((bb[0]+bb[2])/2,bb[1])])
    for a in ['D','E']:
        aa=boxes[a]
        x=(aa[0]+aa[2])/2
        arrow([(375,475),(375,500),(x,500),(x,530)])
        arrow([(x,660),(x,685),(375,685),(375,710)])
    arrow([(690,930),(755,930),(755,95),(820,95)])
    arrow([(690,1130),(790,1130),(790,810),(820,810)])
    path=QA/'architecture.png'
    im.save(path)
    return path


doc = Document()
# Remove inherited title rules from the bundled default Word template.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec=doc.sections[0]
sec.page_width=Inches(8.5)
sec.page_height=Inches(11)
sec.top_margin=Inches(.78)
sec.bottom_margin=Inches(.78)
sec.left_margin=Inches(.85)
sec.right_margin=Inches(.85)
sec.header_distance=Inches(.30)
sec.footer_distance=Inches(.30)

for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
    st=doc.styles[name]
    st.font.name='Calibri'
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.widow_control=True
    if name.startswith('Heading'):
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(14 if name=='Heading 1' else 10)
        st.paragraph_format.space_after=Pt(6)
        st.font.size=Pt({'Heading 1':15,'Heading 2':12.5,'Heading 3':11.5}[name])
    else:
        st.font.size=Pt(11)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(22)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(14)
doc.styles['Title'].paragraph_format.keep_with_next=True
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT

header=sec.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run('OCR pdf V5  •  Đặc tả nghiên cứu'), size=9)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run('Trang '),size=9)
footer._p.append(element('w:fldSimple', **{'w:instr':'PAGE'}))
set_font(footer.add_run(' / '),size=9)
footer._p.append(element('w:fldSimple', **{'w:instr':'NUMPAGES'}))

table_count=0
code_count=0
heading_count=0


def table_widths(headers):
    n=len(headers)
    if n==3:
        if headers[0].strip() in ['Mã','Ca']:
            return [.70,3.30,2.80]
        return [1.40,2.70,2.70]
    if n==4:
        if headers[0].strip() in ['Mã','Cổng','Giai đoạn']:
            return [.72,2.52,2.32,1.24] if headers[0]=='Cổng' else [.72,2.48,2.25,1.35]
        return [1.0,2.40,2.45,.95]
    if n==5:
        return [.88,1.53,2.25,1.40,.74]
    return [6.8/n]*n


def add_table(lines):
    global table_count
    data=[]
    for line in lines:
        cells=[s.strip() for s in re.split(r'(?<!\\)\|', line.strip().strip('|'))]
        if all(re.fullmatch(r':?-{3,}:?',c) for c in cells):
            continue
        data.append(cells)
    headers=data[0]
    widths=table_widths(headers)
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i,w in enumerate(widths):
        t.columns[i].width=Inches(w)
    pr=t._tbl.tblPr
    borders=element('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        borders.append(element('w:'+side, **{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
    pr.append(borders)
    margins=element('w:tblCellMar')
    for side,value in [('top',90),('bottom',90),('left',105),('right',105)]:
        margins.append(element('w:'+side, **{'w:w':value,'w:type':'dxa'}))
    pr.append(margins)
    for row_idx,record in enumerate(data):
        row=t.rows[0] if row_idx==0 else t.add_row()
        row._tr.get_or_add_trPr().append(element('w:cantSplit'))
        if row_idx==0:
            row._tr.get_or_add_trPr().append(element('w:tblHeader', **{'w:val':'true'}))
        for col_idx,cell in enumerate(row.cells):
            cell.width=Inches(widths[col_idx])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill='30475E' if row_idx==0 else ('F2F5F8' if row_idx%2==0 else 'FFFFFF')
            cell._tc.get_or_add_tcPr().append(element('w:shd', **{'w:fill':fill}))
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.08
            p.paragraph_format.keep_with_next=(row_idx==0)
            if widths[col_idx]<1.1:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p, record[col_idx] if col_idx<len(record) else '', table=True, header=row_idx==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)
    table_count+=1


def add_code(lang, lines):
    global code_count
    if lang=='mermaid':
        p=doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(draw_architecture()), width=Inches(6.40))
        cap=doc.add_paragraph('Sơ đồ kiến trúc xử lý và đánh giá với ground truth độc lập')
        cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:
            set_font(r,size=9,italic=True)
        return
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(3)
    p.paragraph_format.space_after=Pt(9)
    p.paragraph_format.line_spacing=1.05
    set_font(p.add_run('\n'.join(lines)), 'Consolas', 9)
    code_count+=1


lines=source_text.splitlines()
i=0
in_final_decisions=False
while i<len(lines):
    s=lines[i].strip()
    if not s or s=='---':
        i+=1
        continue
    if s.startswith('```'):
        lang=s[3:].strip()
        block=[]
        i+=1
        while i<len(lines) and not lines[i].strip().startswith('```'):
            block.append(lines[i])
            i+=1
        add_code(lang,block)
        i+=1
        continue
    if s.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i])
            i+=1
        add_table(block)
        continue
    m=re.match(r'^(#{1,4})\s+(.+)$',s)
    if m:
        depth=len(m[1])
        title=m[2]
        if depth==1:
            p=doc.add_paragraph(style='Title')
            p.add_run('Đặc tả dự án OCR pdf V5').add_break()
            p.add_run('phục vụ nghiên cứu khoa học')
        else:
            if depth==2 and title.startswith('17.'):
                in_final_decisions=True
            title=re.sub(r'[^\w\s]',' ', title, flags=re.UNICODE).replace('_',' ')
            title=re.sub(r'\s+',' ',title).strip()
            doc.add_paragraph(title, style='Heading '+str(depth-1))
            heading_count+=1
        i+=1
        continue
    m=re.match(r'^-\s+(.+)$',s)
    if m:
        p=doc.add_paragraph(style='List Bullet')
        inline(p,m[1])
        i+=1
        continue
    m=re.match(r'^(\d+)\.\s+(.+)$',s)
    if m:
        # Explicit numbering preserves restarts and exact sequence in the source.
        p=doc.add_paragraph()
        p.paragraph_format.left_indent=Inches(.22)
        p.paragraph_format.first_line_indent=Inches(-.22)
        if in_final_decisions and int(m[1])<5:
            p.paragraph_format.keep_with_next=True
        inline(p,m[1]+'. '+m[2])
        i+=1
        continue
    para=[s]
    # Two trailing spaces in Markdown delimit the front-matter metadata lines.
    hard_break=lines[i].endswith('  ')
    i+=1
    while not hard_break and i<len(lines):
        nxt=lines[i].strip()
        if not nxt or nxt.startswith(('#','|','```','- ')) or re.match(r'^\d+\.\s',nxt):
            break
        para.append(nxt)
        hard_break=lines[i].endswith('  ')
        i+=1
    p=doc.add_paragraph()
    inline(p,' '.join(para))

doc.core_properties.title='Đặc tả dự án OCR pdf V5 phục vụ nghiên cứu khoa học'
doc.core_properties.subject='PDF đa dạng sang Markdown với đánh giá khoa học và hiệu năng CPU'
doc.core_properties.author=''
doc.core_properties.keywords='PDF, Markdown, OCR, nghiên cứu, ground truth, CPU'
doc.core_properties.comments='Xuất từ đặc tả Markdown phiên bản 1.0. Không chứa kết quả thực nghiệm mới.'
doc.save(OUTPUT)

# Verify record retention at the OOXML layer before rendering.
all_text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
requirements=set(re.findall(r'\b(?:FR|NF|SCI)-\d{2}\b',source_text))
cases=set(re.findall(r'\bS\d{2}\b',source_text))
missing=[v for v in requirements|cases if v not in all_text]
if missing:
    raise RuntimeError('Missing requirement/case IDs: '+','.join(missing))
report={
    'source':str(SOURCE),'output':str(OUTPUT),
    'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'docx_sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    'tables':table_count,'headings':heading_count,'code_blocks':code_count,
    'requirements_retained':len(requirements),'acceptance_cases_retained':len(cases),
    'source_lines':len(lines),'missing_ids':missing,
}
(QA/'export_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
