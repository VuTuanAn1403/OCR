"""Export the existing PDF-to-keyword report to an editable Word document."""
from pathlib import Path
import re
import subprocess
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.graphics import renderPDF

from build_pdf_keyword_process_doc import register_fonts, diagrams, PARAMETER_TABLES

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/QUY_TRINH_TU_PDF_DEN_TU_KHOA.md'
OUT = ROOT / 'output/word/Quy_trinh_PDF_den_tu_khoa_OCR_pdf.docx'
QA = ROOT / 'tmp/word_keyword_report'
POPPLER = Path(sys.executable).parent.parent / 'native/poppler/Library/bin/pdftoppm.exe'
if not POPPLER.is_file():
    POPPLER = Path(r'C:\Users\ASUS\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe')


def xml(parent, tag, attrs=None):
    e = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        e.set(qn(k), str(v))
    parent.append(e)
    return e


def font(run, name='Arial', size=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    if size:
        run.font.size = Pt(size)


def add_content(p, text, size=None, color=None, force_bold=False):
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`|https?://\S+)', text)
    for token in parts:
        if not token:
            continue
        if token.startswith('http'):
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), p.part.relate_to(token, RT.HYPERLINK, is_external=True))
            r = xml(link, 'w:r')
            pr = xml(r, 'w:rPr')
            xml(pr, 'w:color', {'w:val': '146F77'})
            xml(pr, 'w:rFonts', {'w:ascii': 'Arial', 'w:hAnsi': 'Arial'})
            if size:
                xml(pr, 'w:sz', {'w:val': int(size * 2)})
            xml(r, 'w:t').text = token
            p._p.append(link)
            continue
        bold = token.startswith('**')
        code = token.startswith('`')
        if bold:
            add_content(p, token[2:-2], size=size, color=color, force_bold=True)
            continue
        value = token[2:-2] if bold else token[1:-1] if code else token
        if code and len(value) <= 30:
            value = value.replace(' ', '\u00a0')
        r = p.add_run(value)
        font(r, 'Consolas' if code else 'Arial', (size or 9) if code else size)
        if force_bold:
            r.bold = True
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    register_fonts()
    for key, drawing in diagrams().items():
        renderPDF.drawToFile(drawing, str(QA / f'{key}.pdf'))
        subprocess.run([str(POPPLER), '-r', '240', '-singlefile', '-png',
                        str(QA / f'{key}.pdf'), str(QA / key)], check=True,
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.6)
    sec.left_margin, sec.right_margin = Cm(1.55), Cm(1.55)
    sec.header_distance, sec.footer_distance = Cm(.6), Cm(.7)
    styles = doc.styles
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
        st = styles[name]
        st.font.name = 'Arial'
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial')
    normal = styles['Normal']
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string('172C3B')
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size, shade in [('Title', 22, '000000'), ('Heading 1', 15, '172C3B'),
                               ('Heading 2', 11.8, '146F77')]:
        st = styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(shade)
        st.paragraph_format.space_before = Pt(6 if name != 'Title' else 0)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.keep_with_next = True
        borders = st._element.get_or_add_pPr().find(qn('w:pBdr'))
        if borders is not None:
            borders.getparent().remove(borders)
    for name in ['List Bullet', 'List Number']:
        styles[name].paragraph_format.space_after = Pt(4)
        styles[name].paragraph_format.line_spacing = 1.05
    header = sec.header.paragraphs[0]
    r = header.add_run('OCR_pdf  |  Quy trình từ PDF đến từ khoá  |  V6.3')
    font(r, size=8)
    r.font.color.rgb = RGBColor.from_string('516473')
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('Trang ')
    font(r, size=8)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    doc.core_properties.title = 'Quy trình từ PDF đến kết quả đếm từ khoá'
    doc.core_properties.subject = 'Quy trình kỹ thuật OCR_pdf V6.3'
    doc.core_properties.author = 'OCR_pdf'
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i, figure_key, next_page = 0, None, False
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s == '<pagebreak/>':
            next_page = True
            i += 1
            continue
        if s.startswith('<!-- diagram:'):
            figure_key = s.split(':', 1)[1].split()[0]
            i += 1
            continue
        if s.startswith('```'):
            language = s[3:]
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i])
                i += 1
            i += 1
            if language == 'mermaid':
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(7)
                pic = p.add_run().add_picture(str(QA / f'{figure_key}.png'), width=Cm(17.65))
                pic._inline.docPr.set('descr', f'Sơ đồ quy trình {figure_key}')
            else:
                for line in block:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1
                    p.paragraph_format.keep_with_next = True
                    font(p.add_run(line or ' '), 'Consolas', 8.5)
                p.paragraph_format.keep_with_next = False
                p.paragraph_format.space_after = Pt(6)
            continue
        if s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', c) for c in row):
                    rows.append(row)
                i += 1
            n = len(rows[0])
            fractions = {2: [.32, .68], 3: [.26, .49, .25], 4: [.23, .18, .23, .36]}[n]
            if rows[0][0].startswith('Văn bản'):
                fractions = [.36, .20, .13, .31]
            elif rows[0][0] == 'Môi trường':
                fractions = [.25, .32, .43]
            elif rows[0][0] == 'Nội dung':
                fractions = [.20, .39, .41]
            elif rows[0][0] in PARAMETER_TABLES:
                fractions = PARAMETER_TABLES[rows[0][0]]
            t = doc.add_table(rows=0, cols=n)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for col, f in zip(t.columns, fractions):
                col.width = Cm(17.9 * f)
            pr = t._tbl.tblPr
            borders = xml(pr, 'w:tblBorders')
            for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                xml(borders, 'w:' + side, {'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'})
            margins = xml(pr, 'w:tblCellMar')
            for side, value in [('top', 65), ('bottom', 65), ('left', 90), ('right', 90)]:
                xml(margins, 'w:' + side, {'w:w': value, 'w:type': 'dxa'})
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                trpr = t.rows[-1]._tr.get_or_add_trPr()
                xml(trpr, 'w:cantSplit')
                if ri == 0:
                    xml(trpr, 'w:tblHeader')
                for ci, (cell, text) in enumerate(zip(cells, row)):
                    cell.width = Cm(17.9 * fractions[ci])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    xml(cell._tc.get_or_add_tcPr(), 'w:shd', {'w:fill': '173F54' if ri == 0 else 'F2F5F7' if ri % 2 == 0 else 'FFFFFF'})
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.02
                    if text.isdigit():
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_content(p, text, size=9.2, color='FFFFFF' if ri == 0 else None)
                    if ri == 0:
                        # Keep the header row on the same page as the first data row.
                        p.paragraph_format.keep_with_next = True
                        for r in p.runs:
                            r.bold = True
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1
            font(p.add_run(''), size=2)
            continue
        if s.startswith('# '):
            p = doc.add_paragraph(style='Title')
            add_content(p, s[2:])
        elif s.startswith('## '):
            p = doc.add_paragraph(style='Heading 1')
            add_content(p, s[3:])
            if next_page:
                p.paragraph_format.page_break_before = True
                next_page = False
        elif s.startswith('### '):
            p = doc.add_paragraph(style='Heading 2')
            add_content(p, s[4:])
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_content(p, s[2:])
        elif re.match(r'^\d+\. ', s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.2)
            add_content(p, s)
        else:
            texts = [s]
            while i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].startswith(('#', '|', '```', '<', '- ')):
                if re.match(r'^\d+\. ', lines[i + 1]):
                    break
                i += 1
                texts.append(lines[i].strip())
            content = ' '.join(texts)
            if content.startswith('**Sơ đồ:**'):
                content = ('**Sơ đồ:** ba sơ đồ được vẽ lại từ luồng mã nguồn của dự án. '
                           'Bản Word nhúng sơ đồ dạng hình ảnh; nội dung và bảng có thể chỉnh sửa trực tiếp. '
                           'Bản Markdown giữ mã Mermaid. Công cụ xuất Word là python-docx; '
                           'công cụ này không phải dependency của pipeline OCR.')
            p = doc.add_paragraph()
            add_content(p, content)
        i += 1
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
