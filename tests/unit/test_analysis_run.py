"""P11a: extraction identity is independent of keyword/Excel code, and an
analysis run counts verified extraction without OCR.

The extraction batch below is built with the finalizer test double; it is
never GPU evidence. OCR entry points are replaced by guards that fail the
test if an analysis run reaches them.
"""
import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from src.ocr.application import analysis_run
from src.ocr.application.analysis_run import run_analysis
from src.ocr.application.process_batch import process_batch
from src.ocr.application.process_document import manifest_ocr_hash, ocr_files
from src.ocr.cli.unified import main as cli
from src.ocr.domain.models.processing import IdentityMismatch
from src.ocr.infrastructure.export.completion import mark_completed
from src.ocr.runtime import fingerprint, gpu_scheduler, v63_backend
from src.ocr.runtime.atomic_io import atomic_write_json
from test_v63_pipeline import KEYWORDS, processor_factory, setup_inputs


def tree_hashes(root):
    return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(Path(root).rglob("*")) if p.is_file()}


@pytest.fixture
def extraction(tmp_path):
    source, output = setup_inputs(tmp_path, 3)
    report = process_batch(source, output, {"mode": "test-extraction"}, KEYWORDS, processor_factory([]))
    assert all(r["outcome"] == "SUCCEEDED" for r in report["documents"])
    return output, report


@pytest.fixture
def no_ocr(monkeypatch):
    def guard(*args, **kwargs):
        raise AssertionError("analysis run reached OCR")
    monkeypatch.setattr(gpu_scheduler, "run_document", guard)
    monkeypatch.setattr(v63_backend.GPUProcessor, "__call__", guard)


@pytest.fixture
def counter(monkeypatch):
    calls = []
    original = analysis_run.count_report

    def count(*args, **kwargs):
        calls.append(args[0])
        return original(*args, **kwargs)
    monkeypatch.setattr(analysis_run, "count_report", count)
    return calls


def test_same_dictionary_reproduces_published_counts(extraction, tmp_path, no_ocr):
    source, published = extraction
    before = tree_hashes(source)
    report = run_analysis([source], tmp_path / "analysis", KEYWORDS)
    assert [r["counts"] for r in report["documents"]] == [r["counts"] for r in published["documents"]]
    assert all(r["outcome"] == "SUCCEEDED" for r in report["documents"])
    assert tree_hashes(source) == before  # sources are read-only
    record = json.loads((tmp_path / "analysis/a/source.json").read_text(encoding="utf-8"))
    assert record["extraction"]["kind"] == "extraction" and record["extraction"]["inference_by_this_run"] is False
    assert record["extraction"]["ocr_identity"] == json.loads((source / "a/ocr_receipt.json").read_text(
        encoding="utf-8"))["ocr_identity"]


def test_new_dictionary_counts_without_ocr(extraction, tmp_path, no_ocr, counter):
    source, _ = extraction
    custom = tmp_path / "custom.txt"
    custom.write_text("Internet di động\nChuyển đổi số\n", encoding="utf-8")
    report = run_analysis([source], tmp_path / "analysis", custom, keyword_set_version="custom-1")
    assert counter == ["a", "b", "c"]
    assert [r["counts"] for r in report["documents"]] == [[2, 1]] * 3
    book = openpyxl.load_workbook(report["matrix"], read_only=True)
    assert next(book["Ma trận từ khóa"].iter_rows(values_only=True)) == ("Tệp PDF", "Internet di động", "Chuyển đổi số")
    book.close()


def test_reporting_change_reuses_reports(extraction, tmp_path, monkeypatch, no_ocr, counter):
    source, _ = extraction
    first = tmp_path / "analysis-1"
    run_analysis([source], first, KEYWORDS)
    assert counter == []  # receipts written since P11a record the counting code
    real = fingerprint.code_fingerprints()
    monkeypatch.setattr(fingerprint, "code_fingerprints", lambda: {**real, "publication": "changed-excel-code"})
    counter.clear()
    report = run_analysis([first], tmp_path / "analysis-2", KEYWORDS)
    assert counter == [] and all(r["outcome"] == "SUCCEEDED" for r in report["documents"])
    reused = json.loads((tmp_path / "analysis-2/a/source.json").read_text(encoding="utf-8"))
    assert reused["report_origin"] == "reused" and reused["via_analysis_run"] == [str(first / "a")]
    assert reused["extraction"]["source_dir"] == str(source / "a")


def test_keyword_code_change_recounts_without_ocr(extraction, tmp_path, monkeypatch, no_ocr, counter):
    source, _ = extraction
    first = tmp_path / "analysis-1"
    run_analysis([source], first, KEYWORDS)
    real = fingerprint.code_fingerprints()
    monkeypatch.setattr(fingerprint, "code_fingerprints", lambda: {**real, "analysis": "changed-counter"})
    counter.clear()
    run_analysis([first], tmp_path / "analysis-2", KEYWORDS)
    assert counter == ["a", "b", "c"]


def test_ocr_identity_follows_only_extraction_code():
    request = {"pdf_sha256": "p", "config": {}, "model_hashes": {}, "lock_hashes": {}, "source_fingerprint": "x1",
               "assets": "png", "output_mode": "compact", "text_python": "t", "docling_python": "d",
               "code_fingerprints": {"extraction": "x1", "analysis": "k1", "publication": "r1"}}
    other_analysis = {**request, "code_fingerprints": {"extraction": "x1", "analysis": "k2", "publication": "r2"}}
    other_extraction = {**request, "source_fingerprint": "x2"}
    assert v63_backend.ocr_identity(request) == v63_backend.ocr_identity(other_analysis)
    assert v63_backend.ocr_identity(request) != v63_backend.ocr_identity(other_extraction)


def test_closures_partition_the_code():
    names = lambda closure: {name for name, _ in fingerprint.closure_entries(closure)}
    full, extraction, analysis, publication = (names(c) for c in ("full", "extraction", "analysis", "publication"))
    assert extraction | analysis | publication == full
    assert {path for _, path in fingerprint.closure_entries("full")} == set(fingerprint.closure_files("full"))
    listed = {n for n in full if fingerprint._matches(n, fingerprint.ANALYSIS_CODE + fingerprint.PUBLICATION_CODE)}
    assert listed and not extraction & listed
    for pattern in fingerprint.ANALYSIS_CODE + fingerprint.PUBLICATION_CODE + fingerprint.SHARED_CONTRACTS:
        assert any(fingerprint._matches(name, (pattern,)) for name in full), f"stale closure entry {pattern}"
    keyword_code = "src/ocr_keywords/counter.py"
    assert keyword_code in analysis and keyword_code not in extraction
    assert "src/ocr_workers/phases.py" in extraction
    assert "src/ocr_pipeline/keyword_finalization.py" in analysis - extraction


def test_batch_still_pins_every_closure(monkeypatch):
    backend = {"code_fingerprints": fingerprint.code_fingerprints(), "model_hashes": {}, "lock_hashes": {},
               "source_fingerprint": fingerprint.extraction_fingerprint()}
    v63_backend.verify_runtime(backend)
    real = fingerprint.code_fingerprints()
    monkeypatch.setattr(v63_backend, "code_fingerprints", lambda: {**real, "analysis": "edited"})
    with pytest.raises(IdentityMismatch, match="analysis"):
        v63_backend.verify_runtime(backend)
    legacy = {"model_hashes": {}, "lock_hashes": {}, "source_fingerprint": "pre-p11a"}
    with pytest.raises(IdentityMismatch, match="Source changed"):
        v63_backend.verify_runtime(legacy)


def reseal(directory, page_change=None, manifest_change=None):
    """Forge a self-consistent but invalid source: every hash matches, only the
    evidence is wrong. Verification must reject it on the evidence itself."""
    if page_change:
        path = directory / "pages/page-0001.json"
        page = json.loads(path.read_text(encoding="utf-8"))
        page_change(page)
        atomic_write_json(path, page)
    if manifest_change:
        manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        manifest_change(manifest)
        atomic_write_json(directory / "manifest.json", manifest)
    receipt = json.loads((directory / "ocr_receipt.json").read_text(encoding="utf-8"))
    receipt.update(files=ocr_files(directory), manifest_ocr_sha256=manifest_ocr_hash(directory))
    atomic_write_json(directory / "ocr_receipt.json", receipt)
    mark_completed(directory, "NEEDS_REVIEW", group_validators={"keyword_contract": lambda d, m: None})


@pytest.mark.parametrize("forge,reason", [
    (lambda d: (d / "result.md").write_text("đã sửa", encoding="utf-8"), "Completion marker"),
    (lambda d: reseal(d, page_change=lambda p: p["metadata"]["gpu_ocr"].update(workload="synthetic_probe")),
     "GPU OCR evidence"),
    (lambda d: reseal(d, page_change=lambda p: p["blocks"][0].update(source="native")), "Native text"),
    (lambda d: reseal(d, manifest_change=lambda m: m.update(total_pages=2)), "PARTIAL_DOCUMENT"),
])
def test_invalid_sources_become_failed_documents(extraction, tmp_path, no_ocr, forge, reason):
    source, _ = extraction
    forge(source / "a")
    report = run_analysis([source], tmp_path / "analysis", KEYWORDS)
    by_id = {r["id"]: r for r in report["documents"]}
    assert by_id["a"]["outcome"] == "FAILED_DOCUMENT" and reason in by_id["a"]["error"]
    assert by_id["a"]["counts"] is None
    assert by_id["b"]["outcome"] == by_id["c"]["outcome"] == "SUCCEEDED"


def test_output_must_be_separate_from_sources(extraction):
    source, _ = extraction
    with pytest.raises(ValueError, match="separate"):
        run_analysis([source], source / "analysis", KEYWORDS)


def test_cli_analysis_from(extraction, tmp_path, no_ocr):
    source, _ = extraction
    assert cli(True, ["--analysis-from", str(source), "--output", str(tmp_path / "cli")]) == 0
    assert (tmp_path / "cli/MA_TRAN_DEM_TU_KHOA_3_FILES.xlsx").is_file()
    with pytest.raises(SystemExit):
        cli(True, ["--analysis-from", str(source)])  # an explicit new output is required


def test_publish_failure_then_resume_completes(extraction, tmp_path, monkeypatch, no_ocr):
    from src.ocr.application import publish_keyword_matrix
    source, _ = extraction
    original = publish_keyword_matrix.atomic_write_bytes
    calls = {"n": 0}

    def flaky(path, data):
        calls["n"] += 1
        if calls["n"] == 2:
            raise OSError("matrix locked")
        return original(path, data)
    monkeypatch.setattr(publish_keyword_matrix, "atomic_write_bytes", flaky)
    with pytest.raises(OSError, match="matrix locked"):
        run_analysis([source], tmp_path / "analysis", KEYWORDS)
    report = run_analysis([source], tmp_path / "analysis", KEYWORDS, resume=True)
    assert all(r["outcome"] == "SUCCEEDED" for r in report["documents"])


def test_receipt_without_counting_code_is_recounted(extraction, tmp_path, no_ocr, counter):
    """Pre-P11a receipts carry no analysis fingerprint: count again, never trust."""
    source, _ = extraction
    receipt = json.loads((source / "a/keyword_receipt.json").read_text(encoding="utf-8"))
    receipt.pop("analysis_fingerprint")
    atomic_write_json(source / "a/keyword_receipt.json", receipt)
    mark_completed(source / "a", "NEEDS_REVIEW", group_validators={"keyword_contract": lambda d, m: None})
    run_analysis([source], tmp_path / "analysis", KEYWORDS)
    assert counter == ["a"]


# --- Review F01/F02 regressions (codex review 2026-10-22) ---------------------

def _forged_analysis_source(base, extraction_dir):
    """An analysis-looking source whose provenance is only self-declared."""
    source = base / "forged" / "a"
    source.mkdir(parents=True)
    (source / "result.md").write_text("Chuyển đổi số: arbitrary native text, no OCR receipt", encoding="utf-8")
    sha = hashlib.sha256((source / "result.md").read_bytes()).hexdigest()
    (source / "source.json").write_text(json.dumps({
        "schema_version": "analysis-source-v1", "document_id": "a", "pdf_sha256": "unverified",
        "result_md_sha256": sha, "extraction": {"kind": "extraction", "source_dir": str(extraction_dir),
        "ocr_identity": "unverified", "selected_pages": [1], "total_pages": 1, "validation": "PASS",
        "result_md_sha256": sha, "inference_by_this_run": False}}), encoding="utf-8")
    mark_completed(source, "PASS")
    return source


def test_self_declared_analysis_source_is_rejected(tmp_path, no_ocr):
    source = _forged_analysis_source(tmp_path, tmp_path / "does-not-exist")
    report = run_analysis([source.parent], tmp_path / "analysis", KEYWORDS)
    (row,) = report["documents"]
    assert row["outcome"] == "FAILED_DOCUMENT" and row["counts"] is None
    assert "Extraction evidence" in row["error"]


def test_analysis_source_pointing_at_real_extraction_with_other_text_is_rejected(extraction, tmp_path, no_ocr):
    source_dir, _ = extraction
    source = _forged_analysis_source(tmp_path, source_dir / "a")  # real, verifiable extraction, other text
    report = run_analysis([source.parent], tmp_path / "analysis", KEYWORDS)
    (row,) = report["documents"]
    assert row["outcome"] == "FAILED_DOCUMENT" and "not the result.md of the verified extraction" in row["error"]


def test_edited_chain_text_is_not_legitimized_by_new_hashes(extraction, tmp_path, no_ocr):
    source, _ = extraction
    first = tmp_path / "analysis-1"
    run_analysis([source], first, KEYWORDS)
    doc = first / "a"
    (doc / "result.md").write_text("Internet di động " * 50, encoding="utf-8")
    record = json.loads((doc / "source.json").read_text(encoding="utf-8"))
    record["result_md_sha256"] = hashlib.sha256((doc / "result.md").read_bytes()).hexdigest()
    atomic_write_json(doc / "source.json", record)
    mark_completed(doc, "NEEDS_REVIEW")  # marker re-sealed over the edited text
    report = run_analysis([first], tmp_path / "analysis-2", KEYWORDS)
    by_id = {r["id"]: r for r in report["documents"]}
    assert by_id["a"]["outcome"] == "FAILED_DOCUMENT" and by_id["a"]["counts"] is None
    assert by_id["b"]["outcome"] == by_id["c"]["outcome"] == "SUCCEEDED"


def test_chain_without_its_extraction_is_rejected(extraction, tmp_path, no_ocr):
    import shutil
    source, _ = extraction
    copy = tmp_path / "copy"
    shutil.copytree(source, copy)
    first = tmp_path / "analysis-1"
    run_analysis([copy], first, KEYWORDS)
    shutil.rmtree(copy / "b")  # the extraction behind one chained document disappears
    report = run_analysis([first], tmp_path / "analysis-2", KEYWORDS)
    by_id = {r["id"]: r for r in report["documents"]}
    assert by_id["b"]["outcome"] == "FAILED_DOCUMENT" and "unavailable" in by_id["b"]["error"]
    assert by_id["a"]["outcome"] == by_id["c"]["outcome"] == "SUCCEEDED"


def test_valid_chain_still_succeeds_read_only(extraction, tmp_path, no_ocr):
    source, _ = extraction
    first = tmp_path / "analysis-1"
    run_analysis([source], first, KEYWORDS)
    before = tree_hashes(first), tree_hashes(source)
    report = run_analysis([first], tmp_path / "analysis-2", KEYWORDS)
    assert all(r["outcome"] == "SUCCEEDED" for r in report["documents"])
    assert (tree_hashes(first), tree_hashes(source)) == before


@pytest.mark.parametrize("damage,reason", [
    (lambda d: (d / "result.md").unlink(), "result.md missing"),
    (lambda d: (d / "manifest.json").unlink(), "Completion marker"),
    (lambda d: (d / "ocr_receipt.json").unlink(), "Completion marker"),
    (lambda d: __import__("shutil").rmtree(d), "directory missing"),
])
def test_damaged_batch_document_keeps_its_row(extraction, tmp_path, no_ocr, damage, reason):
    source, _ = extraction
    damage(source / "a")
    report = run_analysis([source], tmp_path / "analysis", KEYWORDS)
    assert [r["id"] for r in report["documents"]] == ["a", "b", "c"]
    by_id = {r["id"]: r for r in report["documents"]}
    assert by_id["a"]["outcome"] == "FAILED_DOCUMENT" and by_id["a"]["counts"] is None
    assert reason in by_id["a"]["error"]
    assert by_id["b"]["outcome"] == by_id["c"]["outcome"] == "SUCCEEDED"
    book = openpyxl.load_workbook(report["matrix"], read_only=True)
    names = [row[0] for row in book["Ma trận từ khóa"].iter_rows(min_row=2, values_only=True)][:-1]
    book.close()
    assert len(names) == 3  # the failed document stays visible in the matrix


def test_broken_standalone_source_is_a_failed_row(tmp_path, no_ocr):
    broken = tmp_path / "standalone"
    broken.mkdir()
    (broken / "notes.txt").write_text("not an output", encoding="utf-8")
    report = run_analysis([broken], tmp_path / "analysis", KEYWORDS)
    (row,) = report["documents"]
    assert row["id"] == "standalone" and row["outcome"] == "FAILED_DOCUMENT" and row["counts"] is None
