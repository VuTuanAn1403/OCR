import unicodedata

from openpyxl import load_workbook
import pytest

from src.ocr.application.count_keywords import CountKeywordsUseCase
from src.ocr.infrastructure.postprocessing.keyword_counter import KeywordCounter


def test_exact_vietnamese_counts_keep_accents_and_ignore_case():
    report = KeywordCounter.count_keywords(
        "TRÍ TUỆ NHÂN TẠO, trí tuệ nhân tạo, tri tue nhan tao. Dữ liệu lớn.",
        ["Trí tuệ nhân tạo", "dữ liệu lớn", "học sâu"], "annual", "result.md")
    assert [(item.keyword, item.exact_count, item.count) for item in report.keywords] == [
        ("Trí tuệ nhân tạo", 2, 2), ("dữ liệu lớn", 1, 1), ("học sâu", 0, 0)]
    assert report.total_keywords_searched == 3
    assert report.total_matches == report.total_exact_matches == 3


def test_unicode_blank_and_duplicate_keywords_do_not_inflate_totals(tmp_path):
    path = tmp_path / "keywords.txt"
    path.write_text("\ufeff  Internet di động \n\nINTERNET DI ĐỘNG\n" +
                    unicodedata.normalize("NFD", "Internet di động") + "\n", encoding="utf-8")
    keywords = KeywordCounter.load_keywords(path)
    report = KeywordCounter.count_keywords(unicodedata.normalize("NFD", "Internet di động"), keywords)
    assert keywords == ["Internet di động"]
    assert report.total_keywords_searched == report.total_matches == 1


def test_literal_substrings_independent_phrases_and_non_overlapping_matches():
    report = KeywordCounter.count_keywords("C++ a.b aXb aaaa khai thác dữ liệu lớn", [
        "C++", "a.b", "aa", "dữ liệu", "dữ liệu lớn", "\n", "aa"])
    assert [item.count for item in report.keywords] == [1, 1, 2, 1, 1]
    assert report.total_matches == 6


def test_exact_matching_does_not_join_whitespace_or_fix_ocr():
    report = KeywordCounter.count_keywords("học\nmáy học  máy học may học máy", ["học máy"])
    assert report.total_matches == 1


def test_empty_keyword_list_and_empty_text():
    assert KeywordCounter.count_keywords("body", []).total_matches == 0
    assert KeywordCounter.count_keywords("", ["keyword"]).keywords[0].count == 0


def test_missing_keyword_file_is_not_created(tmp_path):
    missing = tmp_path / "missing.txt"
    with pytest.raises(FileNotFoundError):
        KeywordCounter.load_keywords(missing)
    assert not missing.exists()


def test_empty_keyword_file_is_reported(tmp_path):
    path = tmp_path / "keywords.txt"
    path.write_text("\n \t\n", encoding="utf-8")
    with pytest.raises(ValueError, match="No keywords"):
        KeywordCounter.load_keywords(path)


def test_excel_report_contents_totals_and_refresh_without_changing_markdown(tmp_path):
    document = tmp_path / "annual"
    document.mkdir()
    source = document / "result.md"
    source.write_text("# Trang 1\nĐỔI MỚI, đổi mới. =SUM và C++", encoding="utf-8")
    before = source.read_bytes()
    keywords = tmp_path / "keywords.txt"
    keywords.write_text("Đổi mới\n=SUM\nC++\nKhông có\n", encoding="utf-8")
    report = CountKeywordsUseCase.execute(document, "annual", keywords)
    assert report.total_matches == 4
    workbook = load_workbook(document / "count_keywords.xlsx")
    try:
        sheet = workbook.active
        assert sheet.title == "Thống kê từ khóa"
        assert sheet["A1"].value == "BÁO CÁO TẦN SUẤT TỪ KHÓA"
        assert [sheet.cell(4, c).value for c in (1, 2, 3)] == [
            "STT", "Từ khóa / Cụm từ", "Số lần xuất hiện"]
        assert [(sheet.cell(r, 2).value, sheet.cell(r, 3).value) for r in range(5, 9)] == [
            ("Đổi mới", 2), ("=SUM", 1), ("C++", 1), ("Không có", 0)]
        assert sheet["B6"].data_type == "s"
        assert sheet["B9"].value == "Tổng cộng"
        assert sheet["C9"].value == 4
    finally:
        workbook.close()
    keywords.write_text("C++\n", encoding="utf-8")
    assert CountKeywordsUseCase.execute(document, "annual", keywords).total_matches == 1
    assert source.read_bytes() == before
    assert not list(document.glob("*.tmp"))


def test_default_keywords_resolve_from_project_when_cwd_changes(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = CountKeywordsUseCase.resolve_keywords_file()
    assert path.is_file()
    assert "Trí tuệ nhân tạo" in KeywordCounter.load_keywords(path)


def test_wheel_uses_bundled_keywords(tmp_path, monkeypatch):
    from ocr_runtime import paths
    path = tmp_path / "keywords.txt"
    path.write_text("học máy", encoding="utf-8")
    monkeypatch.setattr(paths, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(paths, "runtime_resource", lambda name: tmp_path / name)
    assert CountKeywordsUseCase.resolve_keywords_file() == path


def test_another_projects_dictionary_is_never_selected(tmp_path, monkeypatch):
    """F04: an installed wheel run from another Python project (its own
    pyproject.toml and config/keywords.txt) keeps the bundled dictionary."""
    from ocr_runtime import paths
    project = tmp_path / "unrelated"
    (project / "config").mkdir(parents=True)
    (project / "pyproject.toml").write_text('[project]\nname = "unrelated"\n', encoding="utf-8")
    (project / "config/keywords.txt").write_text("không phải từ điển OCR", encoding="utf-8")
    bundled = tmp_path / "bundled" / "keywords.txt"
    bundled.parent.mkdir()
    bundled.write_text("học máy", encoding="utf-8")
    monkeypatch.setattr(paths, "PROJECT_ROOT", project)
    monkeypatch.setattr(paths, "runtime_resource", lambda name: bundled.parent / name)
    assert not paths.is_checkout()
    assert CountKeywordsUseCase.resolve_keywords_file() == bundled
    explicit = tmp_path / "chosen.txt"
    explicit.write_text("dữ liệu", encoding="utf-8")
    assert CountKeywordsUseCase.resolve_keywords_file(explicit) == explicit.resolve()


def test_checkout_uses_its_config_dictionary():
    from ocr_runtime import paths
    assert paths.is_checkout()
    assert CountKeywordsUseCase.resolve_keywords_file() == paths.PROJECT_ROOT / "config/keywords.txt"


def test_missing_result_markdown_is_reported(tmp_path):
    with pytest.raises(FileNotFoundError):
        CountKeywordsUseCase.execute(tmp_path, "missing")
    assert not (tmp_path / "count_keywords.xlsx").exists()
