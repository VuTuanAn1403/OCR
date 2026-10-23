"""V6 runtime contracts: single config source (ENV-13), device selection
(ENV-01), atomic writes and completion markers (ENV-14), cache keys (ENV-04),
batch sharding (ENV-08) and perf aggregation (proposal 6.1/6.5)."""
import json

import pytest

from benchmark.perf.run_perf import (
    CPU_P50_LIMITS, compare_with_previous, page_category, percentile, summarize_pages,
)
from benchmark.baselines.legacy_cli.batch_runner import BatchOCRRunner, parse_shard
from benchmark.baselines.legacy_cli.batch_runner import merge_shard_manifests, shard_report_dir
from benchmark.baselines.legacy_cli.run_ocr import parse_pages_arg, resolve_input_path
from src.ocr.infrastructure.cache.artifact_cache import ArtifactCache
from src.ocr.infrastructure.export.completion import (
    MARKER_NAME, clear_completed, mark_completed, verify_completed,
)
from src.ocr.runtime import config as runtime_config
from src.ocr.runtime.atomic_io import atomic_write_json, atomic_write_text
from src.ocr.runtime.config import ConfigError, config_hash, get_profile, load_profiles
from src.ocr.runtime.device import resolve_device
from src.ocr.runtime import models as runtime_models, run_manifest as runtime_manifest


def _write_profiles(path, body="profiles:\n  quick:\n    render_dpi: 150\n"):
    path.write_text(body, encoding="utf-8")
    return path


def test_profiles_load_from_explicit_path_and_name_is_filled(tmp_path):
    path, profiles = load_profiles(_write_profiles(tmp_path / "p.yaml"))
    assert path == (tmp_path / "p.yaml").resolve()
    assert profiles["quick"] == {"render_dpi": 150, "profile": "quick"}


def test_missing_config_is_an_error_not_a_silent_fallback(tmp_path, monkeypatch):
    monkeypatch.setenv("OCR_CONFIG_PATH", str(tmp_path / "absent.yaml"))
    with pytest.raises(ConfigError, match="OCR_CONFIG_PATH"):
        load_profiles()
    monkeypatch.delenv("OCR_CONFIG_PATH")
    monkeypatch.setattr(runtime_config, "DEFAULT_CONFIG_PATH", tmp_path / "also_absent.yaml")
    with pytest.raises(ConfigError, match="not found"):
        load_profiles()


def test_env_config_path_is_used(tmp_path, monkeypatch):
    monkeypatch.setenv("OCR_CONFIG_PATH", str(_write_profiles(tmp_path / "env.yaml")))
    assert list(load_profiles()[1]) == ["quick"]


@pytest.mark.parametrize("body", ["profiles: {}\n", "other: 1\n", "profiles:\n  bad: 3\n"])
def test_malformed_profiles_are_rejected(tmp_path, body):
    with pytest.raises(ConfigError):
        load_profiles(_write_profiles(tmp_path / "p.yaml", body))


def test_unknown_profile_lists_available_names():
    with pytest.raises(ConfigError, match="Available: quick"):
        get_profile({"quick": {}}, "balanced")


def test_repository_config_defines_only_the_unified_profile():
    _, profiles = load_profiles()
    assert set(profiles) == {"unified_gpu_full_ocr"}


def test_config_hash_is_order_independent_and_sensitive_to_values():
    assert config_hash({"a": 1, "b": {"c": 2}}) == config_hash({"b": {"c": 2}, "a": 1})
    assert config_hash({"a": 1}) != config_hash({"a": 2})


def test_device_resolution(monkeypatch):
    assert resolve_device(None) == "cpu"
    assert resolve_device("cpu") == "cpu"
    monkeypatch.setattr("src.ocr.runtime.device.cuda_available", lambda: False)
    assert resolve_device("auto") == "cpu"
    with pytest.raises(RuntimeError):
        resolve_device("cuda")
    monkeypatch.setattr("src.ocr.runtime.device.cuda_available", lambda: True)
    assert resolve_device("auto") == "cuda"
    with pytest.raises(ValueError):
        resolve_device("tpu")


def test_atomic_write_replaces_without_leaving_temp_files(tmp_path):
    target = tmp_path / "sub" / "out.json"
    atomic_write_json(target, {"v": 1})
    atomic_write_json(target, {"v": 2})
    assert json.loads(target.read_text(encoding="utf-8")) == {"v": 2}
    assert [p.name for p in target.parent.iterdir()] == ["out.json"]


def test_atomic_write_keeps_old_file_when_writing_fails(tmp_path, monkeypatch):
    target = tmp_path / "out.md"
    atomic_write_text(target, "old")

    def broken_replace(src, dst):
        raise OSError("disk full")

    monkeypatch.setattr("src.ocr.runtime.atomic_io.os.replace", broken_replace)
    with pytest.raises(OSError):
        atomic_write_text(target, "new")
    assert target.read_text(encoding="utf-8") == "old"
    assert [p.name for p in tmp_path.iterdir()] == ["out.md"]


def test_completion_marker_detects_missing_and_changed_artifacts(tmp_path):
    (tmp_path / "pages").mkdir()
    (tmp_path / "pages" / "page-0001.md").write_text("a", encoding="utf-8")
    (tmp_path / "result.md").write_text("b", encoding="utf-8")
    assert not verify_completed(tmp_path)
    marker = mark_completed(tmp_path, "PASS")
    assert set(marker["files"]) == {"pages/page-0001.md", "result.md"}
    assert verify_completed(tmp_path)
    # A run manifest written afterwards is not part of the document output.
    (tmp_path / "run_manifest.json").write_text("{}", encoding="utf-8")
    assert verify_completed(tmp_path)
    (tmp_path / "result.md").write_text("changed", encoding="utf-8")
    assert not verify_completed(tmp_path)
    mark_completed(tmp_path, "PASS")
    (tmp_path / "pages" / "page-0001.md").unlink()
    assert not verify_completed(tmp_path)
    clear_completed(tmp_path)
    assert not (tmp_path / MARKER_NAME).exists()


def test_artifact_cache_uses_env_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("OCR_CACHE_DIR", str(tmp_path / "cache"))
    assert ArtifactCache().cache_dir == tmp_path / "cache"
    assert ArtifactCache(str(tmp_path / "explicit")).cache_dir == tmp_path / "explicit"


def test_shards_cover_every_file_exactly_once(tmp_path):
    for name in ("c.pdf", "a.pdf", "e.pdf", "b.pdf", "d.pdf"):
        (tmp_path / name).write_bytes(b"not a pdf")
    seen = []
    for index in range(3):
        runner = BatchOCRRunner.__new__(BatchOCRRunner)
        runner.input_dir, runner.file_list, runner.shard = tmp_path, None, (index, 3)
        runner.force, runner.output_dir = True, tmp_path / "out"
        seen.extend(entry["file_name"] for entry in runner.discover_files())
    assert sorted(seen) == ["a.pdf", "b.pdf", "c.pdf", "d.pdf", "e.pdf"]
    assert len(seen) == len(set(seen))


@pytest.mark.parametrize("value", ["1", "a/b", "3/3", "-1/2", "0/0"])
def test_invalid_shards_are_rejected(value):
    with pytest.raises(SystemExit):
        parse_shard(value)


def test_percentile_interpolates():
    assert percentile([], 50) is None
    assert percentile([4.0], 95) == 4.0
    assert percentile([1.0, 2.0, 3.0, 4.0], 50) == 2.5


def _page(category, seconds, from_cache=False, docling=False, page_type=None):
    return {"category": category, "page_type": page_type or category.upper(), "seconds": seconds,
            "stages": {"primary_ocr": seconds / 2}, "from_cache": from_cache, "docling_called": docling}


def test_page_categories_follow_requirement_groups():
    assert page_category("NATIVE_TEXT") == "native"
    assert page_category("MIXED") == "mixed"
    for page_type in ("SCANNED_IMAGE", "IMAGE_ONLY", "NATIVE_TEXT_LOW_QUALITY", "PRODUCT_SHOWCASE_COLLAGE"):
        assert page_category(page_type) == "scan"


def test_summary_excludes_cached_pages_and_checks_limits():
    pages = [_page("native", 0.1), _page("native", 0.2, docling=True), _page("native", 9.0, from_cache=True),
             _page("scan", 2.0), _page("scan", 3.0), _page("scan", 200.0)]
    summary = summarize_pages(pages, page_timeout=120)
    assert summary["pages_measured"] == 5 and summary["pages_excluded_from_cache"] == 1
    native = summary["categories"]["native"]
    assert native["latency_seconds"]["n"] == 2
    assert native["cpu_p50_limit_seconds"] == CPU_P50_LIMITS["native"]
    assert native["meets_cpu_p50_limit"] is True  # p50 = 0.15
    assert native["latency_seconds_with_docling"]["n"] == 1
    scan = summary["categories"]["scan"]
    assert scan["meets_cpu_p50_limit"] is False
    assert scan["pages_over_timeout"] == 1
    assert scan["perf02_tail_ok"] is False


def test_regression_compare_requires_same_hardware():
    def report(p50, cpu):
        return {"run_manifest": {"hardware": {"processor": cpu}},
                "page_latency": {"categories": {"scan": {"latency_seconds": {"p50": p50}}}}}
    assert compare_with_previous(report(1.2, "x"), report(1.0, "x"))["regressions"][0]["category"] == "scan"
    assert compare_with_previous(report(1.05, "x"), report(1.0, "x"))["regressions"] == []
    other = compare_with_previous(report(5.0, "x"), report(1.0, "y"))
    assert other == {"same_hardware": False, "regressions": []}


@pytest.mark.parametrize("body", ["profiles: [", "profiles:\n  1: {}\n",
                                  "profiles:\n  quick:\n    profile: other\n"])
def test_invalid_yaml_and_profile_identity_are_config_errors(tmp_path, body):
    with pytest.raises(ConfigError):
        load_profiles(_write_profiles(tmp_path / "invalid.yaml", body))


def test_page_selection_preserves_only_explicit_positive_pages():
    assert parse_pages_arg(None) is None
    assert parse_pages_arg(" 1-3,5,3 ") == [1, 2, 3, 5]


@pytest.mark.parametrize("value", ["", "foo", "0", "3-1", "1,", "1-2-3", "1,foo"])
def test_invalid_page_selection_does_not_process_the_whole_document(value):
    import argparse
    with pytest.raises(argparse.ArgumentTypeError):
        parse_pages_arg(value)


def test_cli_finds_pdf_in_persistent_data_directory(tmp_path, monkeypatch):
    persistent = tmp_path / "data"
    persistent.mkdir()
    target = persistent / "annual.pdf"
    target.write_bytes(b"pdf")
    monkeypatch.setenv("OCR_DATA_DIR", str(persistent))
    assert resolve_input_path("annual.pdf") == target.resolve()


def test_models_hash_actual_weights_overrides_and_layout(tmp_path, monkeypatch):
    root, override = tmp_path / "models", tmp_path / "vietocr-external"
    override.mkdir()
    (override / "vgg_seq2seq.pth").write_bytes(b"first")
    docling = root / "docling"
    docling.mkdir(parents=True)
    (docling / "model_manifest.json").write_text('{"expected":"old"}')
    weights = docling / "real_model.safetensors"
    weights.write_bytes(b"actual weights")
    layout = tmp_path / "layout.onnx"
    layout.write_bytes(b"layout weights")
    monkeypatch.setenv("OCR_MODELS_DIR", str(root))
    monkeypatch.setenv("VIETOCR_MODEL_DIR", str(override))
    monkeypatch.setattr(runtime_models, "_rapidocr_models_dir", lambda: None)
    config = {"primary_engine": "vietocr", "docling": {"enabled": True},
              "layout_model": {"model_path": str(layout)}}
    hashes = runtime_models.model_hashes(config)
    assert hashes["vietocr/vgg_seq2seq.pth"] is not None
    assert hashes["docling/real_model.safetensors"] is not None
    assert hashes["layout/model"] is not None
    before = runtime_models.model_fingerprint(config)
    weights.write_bytes(b"changed weights")
    assert runtime_models.model_fingerprint(config) != before
    assert "docling/real_model.safetensors" not in runtime_models.model_hashes({"docling": {"enabled": False}})


def test_untracked_source_contents_change_run_provenance(tmp_path, monkeypatch):
    source = tmp_path / "new_source.py"
    source.write_text("first")
    outputs = {("rev-parse", "HEAD"): b"abc\n", ("status", "--porcelain"): b"?? new_source.py\n",
               ("diff", "HEAD", "--binary"): b"",
               ("ls-files", "--others", "--exclude-standard", "-z"): b"new_source.py\0"}
    monkeypatch.setattr(runtime_manifest, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(runtime_manifest, "_git", lambda *args: outputs[args])
    first = runtime_manifest.git_state()
    source.write_text("second")
    second = runtime_manifest.git_state()
    assert first["uncommitted_diff_sha256"] != second["uncommitted_diff_sha256"]
    assert first["untracked_files_sha256"]["new_source.py"] != second["untracked_files_sha256"]["new_source.py"]


def test_in_progress_run_is_not_reported_as_finished_and_records_cached_docs(tmp_path, monkeypatch):
    monkeypatch.setattr(runtime_manifest, "git_state", lambda: {})
    monkeypatch.setattr(runtime_manifest, "hardware_info", lambda: {})
    monkeypatch.setattr(runtime_manifest, "package_versions", lambda: {})
    monkeypatch.setattr(runtime_manifest, "model_hashes", lambda config: {"weight": "actual"})
    config = {"docling": {"enabled": False}}
    recorder = runtime_manifest.RunRecorder("batch", config_path=tmp_path / "profiles.yaml",
        profile="fast", profile_config=config, device="cpu", cache_dir=tmp_path / "cache",
        resume=True, assets_mode="none")
    config["docling"]["enabled"] = True
    recorder.add_document({"document_id": "cached", "status": "CACHED", "pages_from_cache": [1, 2]})
    path = tmp_path / "run_manifest.json"
    recorder.write(path, final=False)
    checkpoint = json.loads(path.read_text())
    assert checkpoint["status"] == "RUNNING" and checkpoint["finished_at_utc"] is None
    assert checkpoint["config"]["resolved"]["docling"]["enabled"] is False
    assert checkpoint["cache"]["pages_from_cache"] == 2
    recorder.write(path)
    assert json.loads(path.read_text())["finished_at_utc"]


@pytest.mark.parametrize("files", [[], "bad", {"../outside.txt": "bad"}])
def test_corrupt_or_traversing_completion_marker_is_rejected(tmp_path, files):
    (tmp_path / MARKER_NAME).write_text(json.dumps({"schema": "completion-v1",
        "validation_status": "PASS", "files": files}))
    assert verify_completed(tmp_path) is False


def test_completion_marker_detects_unrecorded_artifacts(tmp_path):
    (tmp_path / "result.md").write_text("finished")
    mark_completed(tmp_path, "PASS")
    (tmp_path / "extra.json").write_text("incomplete export")
    assert verify_completed(tmp_path) is False
    with pytest.raises(ValueError):
        mark_completed(tmp_path, "FAIL")


def test_batch_resume_rejects_same_filename_with_changed_pdf_or_model(tmp_path, monkeypatch):
    from src.ocr.domain.models.schema import PAGE_SCHEMA_VERSION
    from src.ocr.runtime.fingerprint import source_fingerprint
    source = tmp_path / "annual.pdf"
    source.write_bytes(b"original")
    runner = BatchOCRRunner.__new__(BatchOCRRunner)
    runner.profile, runner.assets_mode = "fast", "none"
    runner.output_mode = "legacy"
    runner.profile_config = {"profile": "fast"}
    monkeypatch.setattr("benchmark.baselines.legacy_cli.batch_runner.model_fingerprint", lambda config: "model1")
    output = tmp_path / "doc"
    output.mkdir()
    manifest = {"profile": "fast", "configuration": runner.profile_config,
                "page_schema_version": PAGE_SCHEMA_VERSION, "source_fingerprint": source_fingerprint(),
                "assets_mode": "none", "file_hash": ArtifactCache.compute_file_hash(str(source)),
                "models_fingerprint": "model1"}
    (output / "manifest.json").write_text(json.dumps(manifest))
    assert runner._manifest_matches(output, source)
    source.write_bytes(b"replacement")
    assert not runner._manifest_matches(output, source)
    manifest["file_hash"] = ArtifactCache.compute_file_hash(str(source))
    manifest["models_fingerprint"] = "model2"
    (output / "manifest.json").write_text(json.dumps(manifest))
    assert not runner._manifest_matches(output, source)


def _write_shard(output, index, count, entries, *, finished=True):
    directory = shard_report_dir(output, (index, count))
    directory.mkdir(parents=True)
    atomic_write_json(directory / "batch_manifest.json", entries)
    atomic_write_json(directory / "run_manifest.json", {"shard": [index, count],
        "status": "COMPLETED" if finished else "RUNNING", "finished_at_utc": "now" if finished else None,
        "run_id": f"run{index}", "source_fingerprint": "source", "page_schema_version": "schema",
        "config": {"resolved": {"profile": "fast"}}, "models": {"weight": "hash"},
        "device": "cpu", "assets_mode": "none"})


def test_shard_manifests_do_not_overwrite_and_can_be_merged(tmp_path):
    _write_shard(tmp_path, 0, 2, [{"document_id": "b", "file_name": "b.pdf", "status": "COMPLETED"}])
    _write_shard(tmp_path, 1, 2, [{"document_id": "a", "file_name": "a.pdf", "status": "CACHED"}])
    assert shard_report_dir(tmp_path, (0, 2)) != shard_report_dir(tmp_path, (1, 2))
    merged = merge_shard_manifests(tmp_path, 2)
    assert [entry["document_id"] for entry in merged] == ["a", "b"]
    assert json.loads((tmp_path / "batch_manifest.json").read_text()) == merged


@pytest.mark.parametrize("failure", ["unfinished", "duplicate", "different_model"])
def test_shard_merge_rejects_unfinished_or_incompatible_results(tmp_path, failure):
    entry = {"document_id": "a", "status": "COMPLETED"}
    _write_shard(tmp_path, 0, 2, [entry])
    _write_shard(tmp_path, 1, 2, [entry if failure == "duplicate" else {**entry, "document_id": "b"}],
                 finished=failure != "unfinished")
    if failure == "different_model":
        path = shard_report_dir(tmp_path, (1, 2)) / "run_manifest.json"
        run = json.loads(path.read_text())
        run["models"]["weight"] = "other"
        atomic_write_json(path, run)
    with pytest.raises(ValueError):
        merge_shard_manifests(tmp_path, 2)
    assert not (tmp_path / "batch_manifest.json").exists()


def test_batch_timing_excludes_document_and_page_cache_history(tmp_path):
    runner = BatchOCRRunner(input_dir=str(tmp_path / "input"), output_dir=str(tmp_path / "output"),
                            profile="fast", excel_path=None)
    entries = [{"document_id": "current", "file_name": "current.pdf", "status": "COMPLETED",
                "total_pages": 2, "wall_clock_seconds": 2.0, "seconds_per_page": 1.0, "mean_confidence": None},
               {"document_id": "cached", "file_name": "cached.pdf", "status": "CACHED",
                "total_pages": 100, "wall_clock_seconds": 900.0, "seconds_per_page": 9.0, "mean_confidence": 0.0}]
    current = runner.output_dir / "current"
    current.mkdir()
    atomic_write_json(current / "quality.json", {"page_reports": [
        {"page": 1, "from_cache": False, "stage_seconds": {"ocr": 1.0}},
        {"page": 2, "from_cache": True, "stage_seconds": {"ocr": 100.0}}]})
    cached = runner.output_dir / "cached"
    cached.mkdir()
    atomic_write_json(cached / "quality.json", {"page_reports": [
        {"page": 1, "stage_seconds": {"ocr": 1000.0}}]})
    runner.generate_reports(entries, batch_duration=10.0)
    report = json.loads(runner.report_json_path.read_text(encoding="utf-8"))
    assert report["summary"]["average_seconds_per_page"] == 1.0
    assert report["summary"]["pages_reused_from_document_cache"] == 100
    assert report["summary"]["overall_pages_per_minute"] == 12.0
    assert report["bottlenecks"][0]["description"].startswith("1.0s")
    assert [entry["mean_confidence"] for entry in report["files"]] == [None, 0.0]


@pytest.mark.parametrize("output_mode", ["legacy", "compact"])
def test_excel_reports_leave_unavailable_confidence_blank_and_keep_measured_zero(tmp_path, monkeypatch, output_mode):
    from openpyxl import load_workbook
    from scripts import export_7_files_excel, export_batch_excel

    output = tmp_path / "output"
    document = output / "fixture"
    document.mkdir(parents=True)
    atomic_write_json(document / "quality.json", {"page_reports": [
        {"page": p, "page_type": "NATIVE_TEXT", "ocr_confidence": confidence,
         "quality_score": .99, "total_blocks": 1}
        for p, confidence in ((1, None), (2, 0.0))]})
    atomic_write_json(document / "benchmark.json", {"pages_benchmarked": [1, 2], "mean_confidence": None})
    if output_mode == "compact":
        (document / "pages").mkdir()
        for number in (1, 2):
            atomic_write_json(document / f"pages/page-{number:04d}.json", {"page_num": number})
        atomic_write_json(document / "result.json", {"schema_version": "result-index-v1", "selected_pages": [1, 2],
                          "page_refs": [{"page_num": n, "path": f"pages/page-{n:04d}.json"} for n in (1, 2)]})
    else:
        atomic_write_json(document / "result.json", {"pages": [{"page_num": 1}, {"page_num": 2}]})
    atomic_write_json(output / "batch_report.json", {"files": [
        {"document_id": "fixture", "file_name": "fixture.pdf", "status": "COMPLETED", "mean_confidence": None}]})
    monkeypatch.setattr(export_7_files_excel, "PROJECT_ROOT", tmp_path)
    books = [tmp_path / "seven.xlsx", tmp_path / "batch.xlsx"]
    export_7_files_excel.generate_7_files_excel(output, books[0])
    export_batch_excel.build_workbook(output, books[1])
    for path in books:
        workbook = load_workbook(path, read_only=True)
        try:
            pages = workbook["Chi tiet trang"]
            assert pages.cell(2, 6).value is None
            assert pages.cell(3, 6).value == 0.0
            reviews = list(workbook["Can kiem tra"].values)[1:]
            assert len(reviews) == 1 and reviews[0][1] == 2
        finally:
            workbook.close()


def test_batch_retries_layout_model_failure_when_docling_is_disabled(tmp_path, monkeypatch):
    from src.ocr.domain.models.schema import PAGE_SCHEMA_VERSION
    from src.ocr.runtime.fingerprint import source_fingerprint
    runner = BatchOCRRunner.__new__(BatchOCRRunner)
    runner.force, runner.output_dir, runner.profile = False, tmp_path, "fast"
    runner.assets_mode = "none"
    runner.output_mode = "legacy"
    runner.profile_config = {"profile": "fast", "docling": {"enabled": False}}
    monkeypatch.setattr("benchmark.baselines.legacy_cli.batch_runner.model_fingerprint", lambda config: "actual-model")
    doc_dir = tmp_path / "annual"
    (doc_dir / "pages").mkdir(parents=True)
    (doc_dir / "pages/page-0001.md").write_text("Content")
    (doc_dir / "result.md").write_text("Content")
    atomic_write_json(doc_dir / "manifest.json", {"profile": "fast", "page_schema_version": PAGE_SCHEMA_VERSION,
        "source_fingerprint": source_fingerprint(), "assets_mode": "none", "configuration": runner.profile_config,
        "models_fingerprint": "actual-model"})
    for name in ("quality", "benchmark"):
        atomic_write_json(doc_dir / f"{name}.json", {})
    atomic_write_json(doc_dir / "validation.json", {"status": "PASS"})
    for status, cached in (("FAILED_FALLBACK", False), ("MODEL_OUTPUT", True)):
        atomic_write_json(doc_dir / "result.json", {"pages": [{"metadata": {"layout_model": {"status": status}}}]})
        mark_completed(doc_dir, "PASS")
        assert (runner.is_cached("annual", 1) is not None) == cached


def test_vietocr_singleton_reloads_for_changed_weights_config_and_model_directory(tmp_path, monkeypatch):
    from src.ocr.infrastructure.ocr import vietnamese_seq2seq_engine as engine_module
    from unittest.mock import Mock
    cls = engine_module.VietnameseSeq2SeqEngine
    for name, value in (("_instance", None), ("_recognizer", None), ("_detector", None),
                        ("_device", None), ("_model_signature", None)):
        monkeypatch.setattr(cls, name, value)
    first, second = tmp_path / "first", tmp_path / "second"
    for path in (first, second):
        path.mkdir()
        (path / "base.yml").write_text("cnn: vgg\n")
        (path / "vgg-seq2seq.yml").write_text("transformer: seq2seq\n")
        (path / "vgg_seq2seq.pth").write_bytes(b"weights")
    monkeypatch.setenv("VIETOCR_MODEL_DIR", str(first))
    monkeypatch.setattr(engine_module, "RapidOCR", Mock())
    monkeypatch.setattr(engine_module, "Cfg", dict)
    predictor = Mock(side_effect=lambda cfg: object())
    monkeypatch.setattr(engine_module, "Predictor", predictor)
    engine = cls(device="cpu")
    original = engine._recognizer
    cls(device="cpu")
    assert predictor.call_count == 1 and engine._recognizer is original
    (first / "vgg_seq2seq.pth").write_bytes(b"replacement weights")
    cls(device="cpu")
    assert predictor.call_count == 2 and engine._recognizer is not original
    (first / "base.yml").write_text("cnn: vgg\nmodified: true\n")
    cls(device="cpu")
    assert predictor.call_count == 3
    monkeypatch.setenv("VIETOCR_MODEL_DIR", str(second))
    cls(device="cpu")
    assert predictor.call_count == 4
    assert predictor.call_args.args[0]["weights"] == str(second / "vgg_seq2seq.pth")


def test_vietocr_results_keep_recognizer_text_before_cleanup(monkeypatch):
    from src.ocr.infrastructure.ocr import vietnamese_seq2seq_engine as engine_module
    from unittest.mock import Mock
    import numpy as np
    engine = object.__new__(engine_module.VietnameseSeq2SeqEngine)
    engine._detector = Mock(return_value=([[[5, 5], [90, 5], [90, 25], [5, 25]]], None))
    raw = " Doanh thu thua\u0302\u0300n ? dự kiến năm tới năm tới năm tới "
    engine._recognizer = Mock()
    engine._recognizer.predict_batch.return_value = ([raw], [.9])
    result = engine.predict_image(np.zeros((40, 100, 3), dtype=np.uint8))[0]
    assert result["text_raw"] == raw
    assert result["text"] == "Doanh thu thuần - dự kiến năm tới"
    assert result["correction_source"] == "unicode_nfc+trim_whitespace+vietocr_dash_repair+vietocr_repetition_guard"


@pytest.mark.parametrize("raw", ["2025 2025 2025", "Lợi nhuận ? 123 triệu đồng", "1\u00a0234 triệu đồng"])
def test_vietocr_cleanup_cannot_change_protected_numeric_evidence(raw):
    from src.ocr.infrastructure.ocr import vietnamese_seq2seq_engine as engine_module
    from unittest.mock import Mock
    import numpy as np
    engine = object.__new__(engine_module.VietnameseSeq2SeqEngine)
    engine._detector = Mock(return_value=([[[5, 5], [90, 5], [90, 25], [5, 25]]], None))
    engine._recognizer = Mock()
    engine._recognizer.predict_batch.return_value = ([raw], [.9])
    result = engine.predict_image(np.zeros((40, 100, 3), dtype=np.uint8))[0]
    assert result["text_raw"] == result["text"] == raw
    assert result["numeric_cleanup_rejected"] and result["needs_review"]
    assert result["review_reason"] == "numeric_cleanup_rejected"
    assert result["correction_source"] is None
